import re
from collections import Counter
from dataclasses import dataclass
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from wordfreq import zipf_frequency

STOPWORDS=set('a al algo algunas algunos ante antes aunque bajo bien cada como con contra cuando de del desde donde dos el él ella ellas ellos en entre era eran eres es esa ese eso esta estaba estaban estamos están estar este esto fue fueron ha había han hasta hay la las le les lo los más me mi mis mucho muy nada ni no nos nosotros o otra otro para pero poco por porque que qué se sea ser si sí sin sobre son su sus también tan te tiene tienen todo todos tu tú un una uno unos usted ustedes ya y yo'.split())
CEFR_ZIPF={'A1':5.4,'A2':5.0,'B1':4.6,'B2':4.2,'C1':3.8,'C2':0.0}
LEVEL_DESCRIPTIONS={'A1':'Beginner','A2':'Elementary','B1':'Intermediate','B2':'Upper-intermediate','C1':'Advanced','C2':'Proficient'}
NAVY='17365D';BLUE='2F75B5';LIGHT='D9EAF7';PALE='EEF5FB';WHITE='FFFFFF';GREEN='538135';PURPLE='7030A0';ORANGE='C65911'
@dataclass
class AnalysisResult: cefr:str;level_label:str;words:list;phrases:list;stats:dict

def map_duolingo(score):
    return 'A1' if score<=29 else 'A2' if score<=59 else 'B1' if score<=99 else 'B2' if score<=129 else 'C1'
def map_qualification(name):return {'GCSE Foundation / early GCSE':'A1','GCSE Higher / strong GCSE':'A2','AS Level':'B1','A Level':'B2'}[name]
def tokenise(text):return re.findall(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ]+(?:[-'][A-Za-zÁÉÍÓÚÜÑáéíóúüñ]+)?",text.lower())
def sentence_list(text):return [s.strip() for s in re.split(r'(?<=[.!?])\s+|\n+',text) if s.strip()]
def context(word,sents):
    p=re.compile(rf'(?<!\w){re.escape(word)}(?!\w)',re.I)
    return next((s[:240] for s in sents if p.search(s)),'')
def phrases(tokens,limit=15):
    grams=[]
    for n in (2,3):
        for i in range(len(tokens)-n+1):
            g=tokens[i:i+n]
            if all(x in STOPWORDS for x in g) or (g[0] in STOPWORDS and g[-1] in STOPWORDS):continue
            grams.append(' '.join(g))
    ranked=sorted(Counter(grams).items(),key=lambda x:(-x[1],-len(x[0]),x[0]));out=[]
    for p,c in ranked:
        out.append({'phrase':p,'count':c})
        if len(out)>=limit:break
    return out
def analyse(text,cefr,label,count=30):
    toks=tokenise(text);counts=Counter(toks);sents=sentence_list(text);threshold=CEFR_ZIPF[cefr];rows=[]
    for word,n in counts.items():
        if len(word)<3 or word in STOPWORDS:continue
        f=zipf_frequency(word,'es')
        if f>=threshold and cefr!='C2':continue
        rows.append({'word':word,'count':n,'zipf':round(f,2),'difficulty':round(max(0,threshold-f),2),'context':context(word,sents)})
    rows.sort(key=lambda x:(-x['count'],-x['difficulty'],x['word']));rows=rows[:count]
    return AnalysisResult(cefr,label,rows,phrases(toks),{'total_words':len(toks),'unique_words':len(counts),'focus_words':len(rows)})
def reading_profile(text,cefr):
    words=[w for w in tokenise(text) if len(w)>=3 and w not in STOPWORDS]
    if not words:return {'estimated_cefr':'A1','challenge':'Insufficient text','rare_share':0,'average_zipf':0}
    fs=[zipf_frequency(w,'es') for w in words];avg=sum(fs)/len(fs);rare=100*sum(f<CEFR_ZIPF[cefr] for f in fs)/len(fs)
    est='A1' if avg>=5.3 else 'A2' if avg>=4.9 else 'B1' if avg>=4.5 else 'B2' if avg>=4.1 else 'C1' if avg>=3.7 else 'C2'
    order=['A1','A2','B1','B2','C1','C2'];gap=order.index(est)-order.index(cefr);challenge='Comfortable' if gap<0 else 'Suitable stretch' if gap==0 else 'Challenging' if gap==1 else 'Very challenging'
    return {'estimated_cefr':est,'challenge':challenge,'rare_share':round(rare,1),'average_zipf':round(avg,2)}
def fill(c,colour):c.fill=PatternFill('solid',fgColor=colour)
def title(ws,text,end):
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=end);c=ws.cell(1,1,text);c.font=Font(size=20,bold=True,color=WHITE);fill(c,NAVY);c.alignment=Alignment(horizontal='center');ws.sheet_view.showGridLines=False
def head(ws,row,labels,colour=BLUE):
    for i,label in enumerate(labels,1):c=ws.cell(row,i,label);c.font=Font(bold=True,color=WHITE);fill(c,colour)
def build_workbook(result,source_title,text,meta,profile):
    wb=Workbook();ws=wb.active;ws.title='Overview';title(ws,'VOCABULARIO MÁGICO V3 • personalised reading pack',6)
    items=[('Source',source_title or meta.get('filename') or 'Pasted text'),('Learner profile',result.level_label),('Working CEFR',result.cefr),('Input type',meta.get('type')),('Extraction method',meta.get('method')),('OCR confidence',meta.get('ocr_confidence')),('Words analysed',result.stats['total_words']),('Focus words',len(result.words))]
    for r,(a,b) in enumerate(items,3):ws.cell(r,1,a).font=Font(bold=True,color=NAVY);ws.cell(r,2,b)
    ws['D3']='Estimated text band';ws['E3']=profile['estimated_cefr'];ws['D4']='Challenge';ws['E4']=profile['challenge'];ws['D5']='Above learner threshold';ws['E5']=f"{profile['rare_share']}%";ws.column_dimensions['A'].width=25;ws.column_dimensions['B'].width=60;ws.column_dimensions['D'].width=28
    v=wb.create_sheet('Focus vocabulary');title(v,'Focus vocabulary',7);head(v,3,['Spanish word','Your meaning','Times','Frequency','Difficulty','Context sentence','Learned?'])
    for r,x in enumerate(result.words,4):
        for c,val in enumerate([x['word'],'',x['count'],x['zipf'],x['difficulty'],x['context'],''],1):v.cell(r,c,val)
        if r%2==0:
            for c in range(1,8):fill(v.cell(r,c),PALE)
    for c,w in {'A':23,'B':28,'C':12,'D':14,'E':14,'F':75,'G':12}.items():v.column_dimensions[c].width=w
    v.freeze_panes='A4';v.auto_filter.ref=f'A3:G{max(3,v.max_row)}'
    p=wb.create_sheet('Useful phrases');title(p,'Useful phrases and chunks',5);head(p,3,['Phrase','Your meaning','Occurrences','Example / note','Practised?'],PURPLE)
    for r,x in enumerate(result.phrases,4):p.cell(r,1,x['phrase']);p.cell(r,3,x['count'])
    for c,w in {'A':40,'B':30,'C':14,'D':60,'E':14}.items():p.column_dimensions[c].width=w
    q=wb.create_sheet('Quiz');title(q,'Vocabulary quiz',5);head(q,3,['Question','Spanish','Your answer','Correction','Score'],ORANGE)
    for i,x in enumerate(result.words[:15],1):q.cell(3+i,1,f'{i}. What does this mean?');q.cell(3+i,2,x['word'])
    for c,w in {'A':44,'B':24,'C':44,'D':44,'E':12}.items():q.column_dimensions[c].width=w
    d=wb.create_sheet('Discussion');title(d,'Reading and discussion prompts',4);head(d,3,['No.','Prompt','Notes in Spanish','Useful vocabulary'],GREEN)
    prompts=['Resume el texto en tres frases.','¿Cuál es la idea principal?','¿Qué parte te pareció más interesante?','¿Qué palabras nuevas podrías usar?','¿Estás de acuerdo? ¿Por qué?','Cuenta una experiencia relacionada.']
    for i,x in enumerate(prompts,1):d.cell(i+3,1,i);d.cell(i+3,2,x)
    for c,w in {'A':8,'B':58,'C':58,'D':42}.items():d.column_dimensions[c].width=w
    s=wb.create_sheet('Source text');title(s,'Extracted source text',2);s['A3']=text;s['A3'].alignment=Alignment(wrap_text=True,vertical='top');s.column_dimensions['A'].width=120;s.row_dimensions[3].height=420
    m=wb.create_sheet('Method');title(m,'Method and limitations',3);m.append([]);m.append(['Item','Details','Source']);m.append(['OCR','Tesseract with Spanish and English language packs','https://github.com/tesseract-ocr/tesseract']);m.append(['Frequency','Spanish Zipf frequency estimates','https://github.com/rspeer/wordfreq']);m.append(['Caution','OCR and level estimates must be reviewed by the learner',''])
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:cell.alignment=Alignment(horizontal=cell.alignment.horizontal or 'left',vertical=cell.alignment.vertical or 'center',wrap_text=True)
    out=BytesIO();wb.save(out);out.seek(0);return out.getvalue()
